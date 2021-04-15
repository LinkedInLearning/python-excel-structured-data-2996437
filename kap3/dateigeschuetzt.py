from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.chart import PieChart, BarChart, Reference, Series

try:
    wb = load_workbook("daten/datei1.xlsx")
    ws = wb.active
except:
    print("Fehler beim Einlesen der Datei")

values = Reference(ws, min_col=1, min_row=2, max_col=6, max_row=10)
chart = BarChart()
chart.add_data(values)
ws.add_chart(chart, "a20")

values = Reference(ws, min_col=2, min_row=2, max_col=2, max_row=6)
chart = PieChart()
chart.add_data(values)
ws.add_chart(chart, "k20")

ws.protection.sheet = True
ws.protection.enable()

try:
    
    #wb.security.lockStructure = True

    wb.save("daten/dateigeschuetzt.xlsx")
    wb.security.set_workbook_password('geheim')
except Exception:
    print("Fehler beim Schreiben der Exceldatei")
